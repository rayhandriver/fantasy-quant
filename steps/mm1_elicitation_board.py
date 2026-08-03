"""MM-1 — the belief-elicitation workbook: top-150 ADP board as an editable .xlsx.

    uv run --extra reports python steps/mm1_elicitation_board.py [--n 150] [--out PATH]

Session B was originally scoped as ~50 designed pairwise comparisons (see the 2026-08-01
session-5 pointer in CLAUDE.md). The user asked instead for a flat spreadsheet over the top-N
ADP board with a column to record whether he likes each player at his ADP, freeform notes, and
"any other columns of necessary info valuable to creating my personality" — so this pulls every
reference column the frozen engine already carries (market, model, risk, situation) rather than
re-deriving anything, and leaves the belief columns blank for hand entry.

Read-only against the frozen stack: :func:`fantasy_quant.draft.session.build_board` (Phase 4/5
value + distributions), :mod:`fantasy_quant.situation.events` (16.5 event board). Nothing here is
fitted; MM-1's actual model fit (pairwise-comparison + real-pick pooled conditional logit) is a
separate, later step that reads this workbook's "YOUR TAKE" block once it is filled in.

Output: ``reference/mm1_belief_board_<season>_<asof>.xlsx`` (dated like
``coaches_2026_signed_off_*.csv`` — this workbook is the hand-curated artifact, re-generate a
fresh one each time the board moves rather than editing a stale copy in place).
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fantasy_quant.draft import session
from fantasy_quant.situation import events

DB = Path("data/fantasy_quant.duckdb")
SEASON = 2026
TEAMS = 10

# ── column plan ──────────────────────────────────────────────────────────────────────────────
# (group, header, width) — data filled in by _row() below, in the same order.
COLUMNS: list[tuple[str, str, float]] = [
    ("IDENTITY", "Rank (ADP)", 9),
    ("IDENTITY", "Player", 22),
    ("IDENTITY", "Pos", 6),
    ("IDENTITY", "Team", 6),
    ("IDENTITY", "Bye", 6),
    ("IDENTITY", "Rookie", 8),
    ("MARKET", "ADP", 8),
    ("MARKET", "Round (10-tm)", 9),
    ("OUR MODEL", "Proj Pts", 9),
    ("OUR MODEL", "Model Mean", 10),
    ("OUR MODEL", "Games Est /17", 10),
    ("OUR MODEL", "Our Value Rank", 11),
    ("OUR MODEL", "Model vs ADP", 11),
    ("OUR MODEL", "TD Regression", 11),
    ("OUR MODEL", "Role Trend", 10),
    ("RISK PROFILE", "Upside", 8),
    ("RISK PROFILE", "Floor", 8),
    ("RISK PROFILE", "Boom/Bust Spread", 12),
    ("RISK PROFILE", "Boom % (live)", 10),
    ("RISK PROFILE", "Bust % (live)", 10),
    ("RISK PROFILE", "Durability", 10),
    ("SITUATION", "Situation", 14),
    ("SITUATION", "Situation Notes", 44),
    ("YOUR TAKE", "Personal Value", 17),
    ("YOUR TAKE", "Value Score", 10),
    ("YOUR TAKE", "Personal ADP (optional #)", 14),
    ("YOUR TAKE", "Confidence (1-5)", 12),
    ("YOUR TAKE", "Why (short reason)", 30),
    ("YOUR TAKE", "Notes — what you actually believe", 60),
]

# Everything in "YOUR TAKE" is a dropdown or a number except these two — the only free typing.
FREE_TEXT_COLUMNS = {"Why (short reason)", "Notes — what you actually believe"}

# 5-point ordinal: how the player's true value compares to where the market has him. Chosen over
# a "like/dislike" scale because it is directly a signed magnitude — undervalued = a target,
# overvalued = a fade — and it is what "Value Score" below turns into a number without any typing.
VALUE_OPTIONS = ["Very undervalued", "Undervalued", "Evenly valued", "Overvalued", "Very overvalued"]
VALUE_SCORE = {"Very undervalued": 2, "Undervalued": 1, "Evenly valued": 0,
               "Overvalued": -1, "Very overvalued": -2}
CONFIDENCE_OPTIONS = ["1", "2", "3", "4", "5"]  # 1 = pure gut, 5 = would bet on it

VALUE_COLORS = {
    "Very undervalued": "9BC169", "Undervalued": "E2EFDA", "Evenly valued": "FFFFFF",
    "Overvalued": "FCE4D6", "Very overvalued": "FFA5A5",
}

GROUP_FILL = {
    "IDENTITY": "D9D9D9", "MARKET": "DDEBF7", "OUR MODEL": "E2EFDA",
    "RISK PROFILE": "FFF2CC", "SITUATION": "EAD1DC", "YOUR TAKE": "FFF9B1",
}


def _situation_text(row: pd.Series) -> tuple[str, str]:
    """(tag, note) — a short label plus one readable sentence built from the 16.5 event row."""
    if pd.isna(row.get("event_type")):
        return "", ""
    tag = {"team_change": "New team", "new_to_league": "Rookie/new to league",
           "room_change": "Room change", "context_only": "Context"}.get(row["event_type"], row["event_type"])
    bits = []
    if row.get("event_type") == "team_change" and pd.notna(row.get("prev_team")):
        bits.append(f"Moved from {row['prev_team']} to {row['team']}.")
    if row.get("new_play_caller") == 1 and pd.notna(row.get("play_caller")):
        bits.append(f"New play-caller: {row['play_caller']}.")
    if row.get("new_qb") == 1 and pd.notna(row.get("qb")):
        qb = row["qb"] if row["qb"] != "(unsettled)" else "unsettled QB competition"
        bits.append(f"New starting QB: {qb}.")
    if pd.notna(row.get("arrivals")):
        bits.append(f"Arrivals in the room: {row['arrivals']}.")
    if pd.notna(row.get("departures")):
        bits.append(f"Departures: {row['departures']}.")
    if pd.notna(row.get("notes")):
        bits.append(str(row["notes"]))
    return tag, " ".join(bits)


def build_frame(con, season: int, n: int) -> pd.DataFrame:
    res = session.build_board(con, season=season, teams=TEAMS)
    board = res["board"].sort_values("adp").head(n).reset_index(drop=True)
    board["rank_adp"] = range(1, len(board) + 1)

    bye = con.execute(
        """
        select gsis_id, bye
        from ecr_snapshots
        where season = ? and gsis_id is not null
        qualify row_number() over (partition by gsis_id order by as_of desc) = 1
        """,
        [season],
    ).df()
    board = board.merge(bye, on="gsis_id", how="left")

    ev = events.load_events()
    board = board.merge(
        ev, left_on=["name", "team"], right_on=["player", "team"], how="left", suffixes=("", "_ev")
    )

    rows = []
    for _, r in board.iterrows():
        tag, note = _situation_text(r)
        rows.append({
            "Rank (ADP)": int(r["rank_adp"]),
            "Player": r["name"],
            "Pos": r["position"],
            "Team": r["team"],
            "Bye": int(r["bye"]) if pd.notna(r.get("bye")) else None,
            "Rookie": "Yes" if r.get("rookie") == 1.0 else "",
            "ADP": round(float(r["adp"]), 1),
            "Round (10-tm)": math.ceil(float(r["adp"]) / TEAMS),
            "Proj Pts": round(float(r["proj_points"]), 1) if pd.notna(r.get("proj_points")) else None,
            "Model Mean": round(float(r["mean"]), 1) if pd.notna(r.get("mean")) else None,
            "Games Est /17": round(float(r["games_played_mean"]), 1) if pd.notna(r.get("games_played_mean")) else None,
            "Our Value Rank": int(r["overall_rank"]) if pd.notna(r.get("overall_rank")) else None,
            "Model vs ADP": (int(r["rank_adp"]) - int(r["overall_rank"])) if pd.notna(r.get("overall_rank")) else None,
            "TD Regression": round(float(r["td_regression"]), 2) if pd.notna(r.get("td_regression")) else None,
            "Role Trend": round(float(r["role_delta"]), 2) if pd.notna(r.get("role_delta")) else None,
            "Upside": round(float(r["upside"]), 2) if pd.notna(r.get("upside")) else None,
            "Floor": round(float(r["floor"]), 2) if pd.notna(r.get("floor")) else None,
            "Boom/Bust Spread": round(float(r["tail_risk"]), 2) if pd.notna(r.get("tail_risk")) else None,
            "Boom % (live)": round(float(r["boom_prob_live"]), 2) if pd.notna(r.get("boom_prob_live")) else None,
            "Bust % (live)": round(float(r["bust_prob_live"]), 2) if pd.notna(r.get("bust_prob_live")) else None,
            "Durability": round(float(r["durability"]), 2) if pd.notna(r.get("durability")) else None,
            "Situation": tag,
            "Situation Notes": note,
            "Personal Value": None, "Value Score": None, "Personal ADP (optional #)": None,
            "Confidence (1-5)": None, "Why (short reason)": None,
            "Notes — what you actually believe": None,
        })
    return pd.DataFrame(rows)


def write_workbook(df: pd.DataFrame, out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Belief board"

    headers = [h for _, h, _ in COLUMNS]
    groups = [g for g, _, _ in COLUMNS]
    n_cols = len(headers)

    # row 1: merged group bands
    col = 1
    while col <= n_cols:
        g = groups[col - 1]
        span = 1
        while col + span - 1 < n_cols and groups[col + span - 1] == g:
            span += 1
        c1 = get_column_letter(col)
        c2 = get_column_letter(col + span - 1)
        ws.merge_cells(f"{c1}1:{c2}1")
        cell = ws[f"{c1}1"]
        cell.value = g
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill("solid", fgColor=GROUP_FILL[g])
        col += span

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # row 2: column headers
    for i, (group, header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=i, value=header)
        cell.font = Font(bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.fill = PatternFill("solid", fgColor=GROUP_FILL[group])
        cell.border = border
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[2].height = 30

    value_col_letter = get_column_letter(headers.index("Personal Value") + 1)

    # data rows
    for r_i, (_, row) in enumerate(df.iterrows(), start=3):
        for c_i, (group, header, _) in enumerate(COLUMNS, start=1):
            if header == "Value Score":
                ref = f"{value_col_letter}{r_i}"
                conditions = ", ".join(f'{ref}="{k}",{v}' for k, v in VALUE_SCORE.items())
                cell = ws.cell(row=r_i, column=c_i,
                                value=f'=IFS({ref}="","",{conditions},TRUE,"")')
                cell.font = Font(italic=True, color="595959")
                cell.alignment = Alignment(horizontal="center")
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
            else:
                val = row[header]
                cell = ws.cell(row=r_i, column=c_i, value=val if pd.notna(val) else None)
                if group == "YOUR TAKE":
                    cell.fill = PatternFill("solid", fgColor="FFFDE7")
                if header in FREE_TEXT_COLUMNS:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            if header == "Situation Notes":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border

    last_row = 2 + len(df)

    # dropdowns — everything in YOUR TAKE except the two free-text columns and the auto Value Score
    def add_dropdown(header: str, options: list[str]) -> None:
        col_idx = headers.index(header) + 1
        letter = get_column_letter(col_idx)
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                             allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"{letter}3:{letter}{last_row}")

    add_dropdown("Personal Value", VALUE_OPTIONS)
    add_dropdown("Confidence (1-5)", CONFIDENCE_OPTIONS)

    # conditional colour on Personal Value
    from openpyxl.formatting.rule import CellIsRule
    for text, color in VALUE_COLORS.items():
        ws.conditional_formatting.add(
            f"{value_col_letter}3:{value_col_letter}{last_row}",
            CellIsRule(operator="equal", formula=[f'"{text}"'],
                       fill=PatternFill("solid", fgColor=color)),
        )

    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(n_cols)}{last_row}"

    # a short read-me sheet
    notes = wb.create_sheet("Read me")
    notes.column_dimensions["A"].width = 100
    lines = [
        "How to use this workbook",
        "",
        "One row per player, top-150 by current ADP (10-team, full-PPR). Everything left of "
        "\"YOUR TAKE\" is read from the frozen model — market (ADP), our model (projection, "
        "value, availability), risk profile (upside/floor/boom/bust) and situation context "
        "(new team, coaching change, new starting QB). Nothing there needs editing.",
        "",
        "Fill in the yellow \"YOUR TAKE\" columns for as many players as you have an opinion on "
        "— you do not need to do all 150. The ones that matter most for building your personality "
        "are the players where you disagree with the ADP or the model, not the obvious ones.",
        "",
        "The only two columns where you type your own words are \"Why (short reason)\" and "
        "\"Notes\". Everything else in YOUR TAKE is a dropdown or a plain number, so it is quick "
        "to fill in and lands as a clean value with no interpretation needed on the read-back side.",
        "",
        "Personal Value — click the cell for a dropdown. Very undervalued / Undervalued / Evenly "
        "valued / Overvalued / Very overvalued, i.e. how his TRUE value compares to where the "
        "market (ADP) currently has him. Undervalued = a target, overvalued = a fade. Colour-coded "
        "green→red so you can scan a filled sheet at a glance.",
        "Value Score — do not type here. It fills itself in from Personal Value as "
        "+2 / +1 / 0 / −1 / −2, so it is ready to sort, chart or feed a model without re-reading "
        "the words.",
        "Personal ADP (optional #) — only fill this in if you want to be precise about WHERE: the "
        "pick number or round you would actually draft him at. Personal Value already carries the "
        "direction and rough size of your disagreement, so leave this blank unless a number adds "
        "something the dropdown didn't.",
        "Confidence (1-5) — dropdown, 1 = pure gut feel, 5 = you would bet on it. This is about how "
        "sure you are of THIS opinion, not about how good the player is.",
        "Why (short reason) — type a few words for the main driver of your take (e.g. \"new "
        "offense\", \"injury history\", \"model is sleeping on his role\").",
        "Notes — free text, as long as you want. Say exactly what you believe and why — this is "
        "the column the model reads most closely; everything else here is structure around it.",
        "",
        "\"Model vs ADP\" (OUR MODEL, not yours) = your ADP rank minus our value-board rank. "
        "Positive means our model likes him more than the market does (a value by our model); "
        "negative means the market likes him more than our model does. Compare it to your own "
        "Personal Value to see where you and the model agree or split.",
        "",
        f"Generated {pd.Timestamp.now().date()} from the live 2026 board. Re-run "
        "`uv run --extra reports python steps/mm1_elicitation_board.py` for a fresh copy rather "
        "than editing a stale one in place once the board has moved.",
    ]
    for i, line in enumerate(lines, start=1):
        cell = notes.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            cell.font = Font(bold=True, size=13)
    wb.move_sheet("Read me", offset=-1)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=150)
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()

    con = duckdb.connect(str(DB), read_only=True)
    df = build_frame(con, SEASON, args.n)
    asof = pd.Timestamp.now().strftime("%Y%m%d")
    out = args.out or Path(f"reference/mm1_belief_board_{SEASON}_{asof}.xlsx")
    write_workbook(df, out)
    print(f"wrote {len(df)} players -> {out}")


if __name__ == "__main__":
    main()
